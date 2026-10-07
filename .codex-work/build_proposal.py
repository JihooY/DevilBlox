from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


OUT = r"C:\Users\USER\Documents\Github\DevilBlox\2026 AI 창업 경진대회_크리에이터 MOU AI 커뮤니티_아이디어 제안서.docx"
FONT = "맑은 고딕"
NAVY = "17365D"
PALE = "EAF1F8"
LIGHT = "F6F8FA"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, size=10.5, bold=False, color="000000"):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_para(p, before=0, after=5, line=1.28, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep


def add_body(doc, text, bold_lead=None, after=5):
    p = doc.add_paragraph()
    set_para(p, after=after)
    if bold_lead and text.startswith(bold_lead):
        style_run(p.add_run(bold_lead), bold=True)
        style_run(p.add_run(text[len(bold_lead):]))
    else:
        style_run(p.add_run(text))
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_para(p, after=3)
    p.paragraph_format.left_indent = Cm(0.45 + level * 0.4)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    style_run(p.add_run(text), size=10.2)
    return p


def add_heading(doc, number, title):
    p = doc.add_paragraph(style="Heading 1")
    set_para(p, before=0, after=8, line=1.0, keep=True)
    style_run(p.add_run(f"{number}. {title}"), size=15, bold=True, color="000000")
    return p


def add_subheading(doc, title):
    p = doc.add_paragraph(style="Heading 2")
    set_para(p, before=7, after=4, line=1.0, keep=True)
    style_run(p.add_run(title), size=11.5, bold=True, color="000000")
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=0, line=1.05)
        style_run(p.add_run(h), size=font_size, bold=True, color="FFFFFF")
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2 == 1:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_para(p, after=0, line=1.15)
            style_run(p.add_run(value), size=font_size, bold=(i == 0))
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

styles = doc.styles
styles["Normal"].font.name = FONT
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
styles["Normal"].font.size = Pt(10.5)
for style_name in ("Title", "Heading 1", "Heading 2"):
    styles[style_name].font.name = FONT
    styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

# Remove Word's theme border from the built-in Title style.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(78)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_run(p.add_run("2026 AI 창업 경진대회"), size=18, bold=True)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para(p, before=18, after=12, line=1.2)
style_run(p.add_run("아이디어 제안서"), size=28, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para(p, before=22, after=2)
style_run(p.add_run("버츄얼 AI 아바타를 결합한"), size=16, bold=True, color=NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para(p, after=26)
style_run(p.add_run("크리에이터 MOU 커뮤니티 구축"), size=16, bold=True, color=NAVY)

cover = doc.add_table(rows=3, cols=2)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.autofit = False
set_table_borders(cover)
cover_data = [("아이디어명", "크리에이터 MOU AI 커뮤니티"), ("팀명", "[기입]"), ("팀장명", "[기입]")]
for r, (label, value) in zip(cover.rows, cover_data):
    r.cells[0].width = Cm(3.2)
    r.cells[1].width = Cm(10.5)
    set_cell_shading(r.cells[0], PALE)
    for idx, text in enumerate((label, value)):
        c = r.cells[idx]
        set_cell_margins(c, top=150, bottom=150)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
        set_para(p, after=0)
        style_run(p.add_run(text), size=11, bold=(idx == 0))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para(p, before=36)
style_run(p.add_run("제출 전 팀명과 팀장명을 반드시 입력해 주십시오"), size=9, color="666666")

doc.add_page_break()

# Page 1
add_heading(doc, "1", "아이디어의 제안 배경")
add_subheading(doc, "시장과 운영 현장의 문제")
add_body(doc, "크리에이터 커뮤니티는 방송 플랫폼, 영상 플랫폼, 팬카페, SNS 등 여러 채널에 분산되어 있다. 공지와 일정이 채널별로 반복 게시되고, 신규 팬은 필요한 정보를 찾기 어렵다. Discord를 중심 허브로 사용하더라도 운영진이 상시 질문에 대응하고 이벤트를 진행하려면 상당한 반복 업무가 발생한다.")
add_body(doc, "버츄얼 크리에이터와 팬의 관계에서는 말투와 세계관의 일관성이 중요하지만, 크리에이터 본인이 24시간 응답하는 방식은 현실적이지 않다. 반대로 AI가 크리에이터를 대신해 자유롭게 대화하면 오발언, 관계 오인, 브랜드 훼손, 소속사의 통제 상실에 대한 반감이 커질 수 있다.")

add_subheading(doc, "제안 목적")
add_body(doc, "본 아이디어는 크리에이터의 목소리·말투·성격·세계관을 반영한 버츄얼 AI 아바타를 Discord 통합 커뮤니티에 배치하되, 완전 자율 대화형 대리인이 아니라 통제 가능한 운영 보조 에이전트로 설계한다. AI는 공지 안내, 일정 알림, 자주 묻는 질문, 콘텐츠 탐색, 이벤트 진행, 가벼운 정형 소통을 맡고, 민감하거나 관계성이 필요한 대화는 사람 운영진에게 연결한다.")

add_table(doc, ["이해관계자", "현재 불편", "제안 가치"], [
    ("크리에이터", "방송 외 시간의 공지·반복 문의 대응 부담", "본인의 표현 자산을 활용한 제한적 팬 접점 확대"),
    ("소속사·운영진", "채널 분산, 정책 위반 위험, 이벤트 운영 인력 부족", "승인된 지식과 행동 범위 안에서 운영 자동화"),
    ("팬·커뮤니티", "정보 탐색의 어려움, 참여 기회의 편차", "한곳에서 일정·공지·콘텐츠·이벤트에 접근"),
], widths=[3.0, 6.1, 7.3])

add_subheading(doc, "사회적 의미")
add_body(doc, "소규모 크리에이터도 대형 소속사 수준의 커뮤니티 운영 기능을 사용할 수 있어 팬 관리의 격차를 줄일 수 있다. 또한 사람의 정체성을 무단 복제하는 AI가 아니라, 당사자와 소속사의 동의·승인·회수 절차를 전제로 한 퍼블리시티 권리 친화형 활용 모델을 제시한다.")

doc.add_page_break()

# Page 2
add_heading(doc, "2", "아이디어의 소개 및 차별점")
add_subheading(doc, "서비스 개요")
add_body(doc, "가칭 ‘MOU AI 커뮤니티’는 여러 플랫폼의 공지와 콘텐츠를 Discord로 모으고, 크리에이터별 AI 아바타가 승인된 범위에서 운영을 보조하는 B2B2C 커뮤니티 솔루션이다. MOU는 크리에이터·소속사·플랫폼 운영자 간 사용 권한, 데이터 범위, 수익 배분, 금지 행동을 명확히 합의하는 운영 계약 체계를 뜻한다.")

add_table(doc, ["기능", "구현 내용", "사람의 통제"], [
    ("통합 허브", "방송·영상·SNS의 일정, 새 콘텐츠, 공지를 Discord 채널별로 정리", "연동 플랫폼과 게시 규칙을 운영자가 선택"),
    ("페르소나 AI", "승인된 말투·성격·세계관 및 선택적으로 합성 음성을 적용", "학습 자료, 금칙어, 답변 범위, 음성 사용을 당사자가 승인"),
    ("운영 보조", "FAQ, 일정 안내, 콘텐츠 추천, 신규 가입 안내, 문의 분류", "불확실하거나 민감한 문의는 답변하지 않고 담당자에게 이관"),
    ("이벤트 진행", "퀴즈, 출석, 팬미션, 투표, 방송 연계 보상 이벤트 운영", "이벤트 문안·기간·보상은 사전 승인 후 실행"),
    ("안전 관리", "로그 기록, 신고 접수, 위험 표현 탐지, 즉시 중지 기능", "관리자 대시보드에서 수정·회수·중단 가능"),
], widths=[2.5, 8.0, 5.9], font_size=8.7)

add_subheading(doc, "핵심 차별점")
add_bullet(doc, "크리에이터를 대체하는 무제한 자율 대화가 아니라, 운영 보조·가벼운 소통·이벤트에 초점을 둔 제한형 에이전트")
add_bullet(doc, "AI임을 항상 표시하고, 실제 크리에이터의 발언으로 오해될 수 있는 표현과 사적 관계 형성을 제한")
add_bullet(doc, "크리에이터별 권한과 금지 항목을 MOU로 정리해 음성·페르소나·데이터 사용의 불확실성을 낮춤")
add_bullet(doc, "단일 챗봇이 아니라 여러 크리에이터가 입점하고 합동 이벤트를 열 수 있는 커뮤니티 네트워크 구조")

add_subheading(doc, "안전 설계 원칙")
add_body(doc, "정치·의료·법률·투자 조언, 개인정보 요청, 사적 만남 약속, 공격적 논쟁, 미승인 광고 등은 응답 범위에서 제외한다. 모든 발언에는 AI 표식을 붙이고, 주요 이벤트와 공지는 사람의 사전 승인을 거친다. 크리에이터는 언제든 음성 및 페르소나 사용을 중지하거나 자료 삭제를 요청할 수 있다.")

doc.add_page_break()

# Page 3
add_heading(doc, "3", "아이디어의 실현가능성 및 사업성")
add_subheading(doc, "기술 구성과 구현 방식")
add_table(doc, ["구성", "적용 기술", "실현 방법"], [
    ("플랫폼 연동", "Discord Bot, 플랫폼 API, Webhook", "공지·일정·콘텐츠 메타데이터를 수집해 중복을 제거하고 채널별 게시"),
    ("지식 응답", "검색증강생성 RAG, 벡터 검색", "승인된 FAQ·공지·콘텐츠 설명만 검색해 근거 기반으로 답변"),
    ("페르소나", "프롬프트 정책, 예시 대화, 선택형 TTS", "말투는 표현 규칙으로 제한하고 음성은 별도 동의가 있는 경우에만 사용"),
    ("안전·감사", "콘텐츠 필터, 권한 관리, 로그, 관리자 승인", "위험 주제 차단, 신뢰도 임계치 미달 시 이관, 발언 이력 추적"),
], widths=[2.6, 5.4, 8.4], font_size=8.8)

add_subheading(doc, "단계별 실증 계획")
add_table(doc, ["단계", "기간", "주요 산출물", "검증 목표"], [
    ("1단계", "1~2개월", "Discord 통합 봇, 공지·FAQ, 관리자 설정", "공지 게시 성공률 95% 이상, 금지 주제 차단 시나리오 통과"),
    ("2단계", "3~4개월", "크리에이터 1~3팀 페르소나, 이벤트 기능, 로그 대시보드", "반복 문의 자동 처리율 50% 이상, 운영자 개입 시간 30% 절감"),
    ("3단계", "5~6개월", "다중 크리에이터 커뮤니티, 합동 이벤트, 유료 플랜", "월간 재방문율·이벤트 참여율·유료 전환 의향 검증"),
], widths=[2.0, 2.5, 6.4, 5.5], font_size=8.5)

add_subheading(doc, "수익 구조")
add_bullet(doc, "소속사·크리에이터 대상 월 구독: 연동 채널 수, 커뮤니티 규모, AI 사용량, 관리자 기능에 따른 요금제")
add_bullet(doc, "온보딩·커스터마이징 비용: 페르소나 가이드, 안전 정책, 지식베이스, 이벤트 템플릿 구축")
add_bullet(doc, "합동 이벤트·브랜드 캠페인 수수료: 참여형 이벤트 운영과 성과 리포트 제공")
add_bullet(doc, "향후 팬 멤버십 및 디지털 상품 연계 수수료: 운영 주체와 사전 합의된 항목에 한해 적용")

add_subheading(doc, "운영 리스크와 대응")
add_body(doc, "가장 큰 리스크는 페르소나 오남용, 환각 답변, 음성권 침해, 팬의 관계 오인이다. 이를 줄이기 위해 승인 자료 기반 답변, 금지 주제 필터, 답변 신뢰도 기준, 관리자 즉시 중지, 로그 감사, AI 표시, 데이터 삭제 절차를 기본 기능으로 제공한다. 공개 전에는 크리에이터 측 검수 시나리오를 통과한 기능만 활성화한다.")

doc.add_page_break()

# Page 4
add_heading(doc, "4", "아이디어의 기대 효과")
add_subheading(doc, "정량적 목표")
add_table(doc, ["영역", "시범 운영 목표", "측정 방법"], [
    ("운영 효율", "반복 문의 대응 시간 30% 이상 절감", "운영 전후 문의 처리 시간 및 담당자 투입 시간 비교"),
    ("정보 접근", "공지·일정 자동 게시 성공률 95% 이상", "연동 이벤트 대비 정상 게시 건수 집계"),
    ("팬 참여", "이벤트 참여율 및 월간 재방문율 10% 이상 개선", "도입 전 기준 기간과 시범 운영 기간 비교"),
    ("안전성", "중대 정책 위반 발언 0건 목표", "자동 필터·신고·관리자 검수 로그 점검"),
    ("사업성", "시범 참여팀의 유료 전환 의향 50% 이상", "종료 인터뷰 및 가격 수용도 조사"),
], widths=[2.7, 7.1, 6.6], font_size=8.8)

add_subheading(doc, "기대되는 변화")
add_bullet(doc, "크리에이터는 반복 운영 업무를 줄이면서도 방송 외 시간의 팬 접점을 안정적으로 유지할 수 있다.")
add_bullet(doc, "팬은 흩어진 공지와 콘텐츠를 한곳에서 확인하고, 시간대와 관계없이 기본 안내와 이벤트에 참여할 수 있다.")
add_bullet(doc, "소속사는 AI의 행동 범위와 음성·페르소나 사용 조건을 계약과 시스템 권한으로 관리할 수 있다.")
add_bullet(doc, "여러 크리에이터가 합동 이벤트와 교차 홍보를 진행해 개별 팬덤을 넘어선 협업 커뮤니티를 형성할 수 있다.")
add_bullet(doc, "운영·콘텐츠 기획·AI 안전 검수·커뮤니티 매니지먼트 등 새로운 직무 수요를 만들 수 있다.")

add_subheading(doc, "확장 가능성")
add_body(doc, "초기에는 Discord 기반 커뮤니티 운영 보조에 집중하고, 검증 후 팬카페·웹 커뮤니티·라이브 방송 채팅으로 연동 범위를 확장한다. 버츄얼 크리에이터뿐 아니라 교육자, 게임 길드, 아티스트, 브랜드 앰배서더 등 정체성과 커뮤니티 운영이 결합된 분야에도 적용할 수 있다.")

add_subheading(doc, "최종 목표")
add_body(doc, "크리에이터의 개성과 권리를 보호하면서 팬 커뮤니티의 정보 접근성과 참여 경험을 높이는 ‘사람 중심 AI 운영 표준’을 만드는 것이 목표다. 성공 여부는 대화량 자체가 아니라 운영 시간 절감, 정보 전달 정확도, 이벤트 참여, 안전사고 발생 여부, 크리에이터와 팬의 만족도로 평가한다.")

# Footer with page field-like static label; page numbers omitted to avoid stale fields.
for sec in doc.sections:
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(fp, after=0)
    style_run(fp.add_run("2026 AI 창업 경진대회 아이디어 제안서"), size=8, color="777777")

doc.core_properties.title = "버츄얼 AI 아바타를 결합한 크리에이터 MOU 커뮤니티 구축"
doc.core_properties.subject = "2026 AI 창업 경진대회 아이디어 제안서"
doc.core_properties.keywords = "AI 아바타, 크리에이터, Discord, 커뮤니티, MOU"
doc.save(OUT)
print(OUT)
